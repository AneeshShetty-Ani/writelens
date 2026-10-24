"""Extract plain prose from uploaded documents (.docx, .pdf, .tex, .md, .txt) for analysis."""

import io
import re
import zipfile

SUPPORTED = ("docx", "pdf", "tex", "md", "txt")
MAX_BYTES = 10 * 1024 * 1024


class DocumentError(Exception):
    pass


def extract_text(filename: str, data: bytes) -> str:
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if ext not in SUPPORTED:
        raise DocumentError(f"Unsupported file type '.{ext}'. Use: {', '.join(SUPPORTED)}.")
    if len(data) > MAX_BYTES:
        raise DocumentError(f"File is too large (max {MAX_BYTES // (1024 * 1024)} MB).")

    if ext == "docx":
        text = docx_to_text(data)
    elif ext == "pdf":
        text = pdf_to_text(data)
    else:
        raw = data.decode("utf-8", errors="replace")
        text = {"tex": tex_to_text, "md": md_to_text, "txt": lambda s: s}[ext](raw)

    text = tidy(text)
    if not text:
        raise DocumentError("No readable text found in this file.")
    return text


def tidy(text: str) -> str:
    text = text.replace("\r\n", "\n")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r" *\n *", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def docx_to_text(data: bytes) -> str:
    import docx

    try:
        document = docx.Document(io.BytesIO(data))
    except (zipfile.BadZipFile, KeyError, ValueError) as exc:
        raise DocumentError("This doesn't look like a valid .docx file.") from exc

    parts = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    parts.append(cell.text)
    return "\n\n".join(parts)


def pdf_to_text(data: bytes) -> str:
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError

    try:
        reader = PdfReader(io.BytesIO(data))
        # many PDFs are "encrypted" with an empty password only to block editing; those still open
        if reader.is_encrypted and not reader.decrypt(""):
            raise DocumentError("This PDF is password-protected. Remove the password and upload it again.")
        pages = [page.extract_text() or "" for page in reader.pages]
    except (PdfReadError, ValueError, KeyError) as exc:
        raise DocumentError("This doesn't look like a valid .pdf file.") from exc

    text = "\n\n".join(pages)
    # PDFs break lines at the page edge: rejoin hyphenated words, then lines inside a paragraph.
    text = re.sub(r"(\w)-\n(\w)", r"\1\2", text)
    text = re.sub(r"(?<![.!?:\n])\n(?!\n)", " ", text)
    if not text.strip():
        raise DocumentError("No text found in this PDF. If it's a scan (a picture of pages), it has no text to read.")
    return text


# LaTeX commands whose argument is prose worth analysing.
TEX_KEEP_ARG = r"(?:textbf|textit|emph|underline|text|section|subsection|subsubsection|paragraph|chapter|title|caption|footnote)\*?"
# LaTeX commands whose argument is not prose (references, files, labels).
TEX_DROP_ARG = r"(?:cite\w*|ref|eqref|autoref|cref|label|includegraphics|input|include|usepackage|documentclass|bibliography\w*|url|href|hspace|vspace)\*?"
TEX_MATH_ENVS = r"(?:equation|align|gather|multline|eqnarray|math|displaymath|tabular|verbatim|lstlisting|figure|table)\*?"


def tex_to_text(src: str) -> str:
    src = re.sub(r"(?<!\\)%.*", "", src)  # comments

    body = re.search(r"\\begin\{document\}(.*?)(?:\\end\{document\}|$)", src, re.S)
    if body:
        src = body.group(1)

    src = re.sub(rf"\\begin\{{({TEX_MATH_ENVS})\}}.*?\\end\{{\1\}}", " ", src, flags=re.S)
    src = re.sub(r"\$\$.*?\$\$|\\\[.*?\\\]|\\\(.*?\\\)", " ", src, flags=re.S)
    src = re.sub(r"(?<!\\)\$.*?(?<!\\)\$", " ", src, flags=re.S)

    src = re.sub(rf"\\{TEX_DROP_ARG}(?:\[[^\]]*\])*(?:\{{[^}}]*\}})?", "", src)
    # Headings go on their own line so they aren't merged into the next sentence.
    src = re.sub(r"\\(?:section|subsection|subsubsection|chapter|paragraph)\*?\{([^}]*)\}", r"\n\n\1\n\n", src)
    src = re.sub(rf"\\{TEX_KEEP_ARG}\{{([^}}]*)\}}", r"\1", src)
    src = re.sub(r"\\(?:begin|end)\{[^}]*\}", "\n", src)
    src = re.sub(r"\\item\b", "\n", src)

    src = src.replace("\\\\", "\n").replace("~", " ").replace("``", '"').replace("''", '"')
    src = re.sub(r"\\([%$&#_{}])", r"\1", src)  # escaped characters
    src = re.sub(r"\\[A-Za-z]+\*?(?:\[[^\]]*\])?", "", src)  # any remaining commands
    src = src.replace("{", "").replace("}", "")
    return src


def md_to_text(src: str) -> str:
    src = re.sub(r"```.*?```|~~~.*?~~~", "\n", src, flags=re.S)  # code blocks
    src = re.sub(r"`[^`]*`", "", src)  # inline code
    src = re.sub(r"!\[[^\]]*\]\([^)]*\)", "", src)  # images
    src = re.sub(r"\[([^\]]*)\]\([^)]*\)", r"\1", src)  # links -> link text
    src = re.sub(r"<[^>]+>", "", src)  # HTML tags
    src = re.sub(r"^\s{0,3}#{1,6}\s*", "\n", src, flags=re.M)  # headings
    src = re.sub(r"^\s*(?:[-*+]|\d+\.)\s+", "", src, flags=re.M)  # list markers
    src = re.sub(r"^\s*>\s?", "", src, flags=re.M)  # blockquotes
    src = re.sub(r"(\*\*|__|\*|_)(\S.*?\S|\S)\1", r"\2", src)  # bold / italic
    return src
