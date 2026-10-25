"""Build a downloadable report (Markdown, Word or PDF) from the analysis results."""

import io
from datetime import date
from pathlib import Path

from analysis import ENDING_ADVICE, LONG_SENTENCE_WORDS, READ_WPM, SPEAK_WPM, Report, duration, expand_contraction
from nlp_checks import REPEAT_MIN, NlpReport
from score import score_label, writing_score

TITLE = "WriteLens report"
NO_SPACY = "Grammar checks were unavailable (spaCy model not installed)."


def counted(items: list[tuple[str, int]]) -> str:
    return ", ".join(f"{w} ×{n}" for w, n in items)


def stats_rows(report: Report, nlp: NlpReport) -> list[tuple[str, str]]:
    score = writing_score(report, nlp)
    return [
        ("Writing score", f"{score} / 100 ({score_label(score)})"),
        ("Words", str(report.words)),
        ("Sentences", str(report.sentences)),
        ("Avg sentence length", f"{report.avg_sentence_len:.1f}"),
        ("Reading ease (Flesch)", f"{report.reading_ease:.0f}"),
        ("Grade level (Flesch-Kincaid)", f"{report.grade_level:.0f}"),
        ("Reading time", duration(report.words, READ_WPM)),
        ("Speaking time", duration(report.words, SPEAK_WPM)),
    ]


def sections(report: Report, nlp: NlpReport) -> list[tuple[str, list[str]]]:
    """(title, bullet points) for each kind of issue that was found."""
    found = [
        ("Weak opening", [f"“{p}” → open with your own point instead" for p in report.weak_openings]),
        ("Conclusion", [ENDING_ADVICE[k][0].upper() + ENDING_ADVICE[k][1:] for k in report.ending_issues]),
        ("Spelling", [f"{w} → {fix}" + (f" ×{n}" if n > 1 else "") for w, fix, n in report.misspellings]),
        ("Wordy phrases", [f"“{p}” ×{n} → consider {s}" for p, s, n in report.wordy]),
        ("Clichés", [f"“{c}” ×{n} → say it in your own words" for c, n in report.cliches]),
        (f"Long sentences (over {LONG_SENTENCE_WORDS} words)", list(report.long_sentences)),
        ("Passive voice", list(nlp.passive)),
        (f"Repeated words ({REPEAT_MIN}+ uses)", [f"{w} ×{n}" for w, n in nlp.repeated]),
        ("Weak verbs", [f"{w} ×{n}" for w, n in nlp.weak_verbs]),
        ("Filler adverbs", [f"{w} ×{n}" for w, n in nlp.adverbs]),
        ("Contractions (avoid in formal writing)", [f"{w} ×{n} → {expand_contraction(w)}" for w, n in report.contractions]),
        ("Informal words", [f"“{p}” ×{n} → {s}" for p, s, n in report.informal]),
        ("Talking to the reader",
         [f"“you” / “your” ×{report.second_person} → try “people”, “students” or “one”"] if report.second_person else []),
        ("Repeated sentence starts",
         [f"{n} sentences in a row begin with “{w.capitalize()}”" for w, _, n in report.repeated_starts]),
    ]
    if report.uk_spelling:
        found.append(("Mixed UK/US spelling", [
            f"UK: {counted(report.uk_spelling)}",
            f"US: {counted(report.us_spelling)}",
            f"Pick one and use it throughout; the fewest changes: {counted(report.minority_spelling)}",
        ]))
    return [(title, items) for title, items in found if items]


def build_markdown(report: Report, nlp: NlpReport, source: str = "Pasted text") -> str:
    lines = [
        f"# {TITLE}",
        "",
        f"**Source:** {source}  ",
        f"**Date:** {date.today().isoformat()}",
        "",
        "## Stats",
        "",
        "| Metric | Value |",
        "|---|---|",
        *(f"| {name} | {value} |" for name, value in stats_rows(report, nlp)),
        "",
        "## Suggestions",
        "",
    ]
    found = sections(report, nlp)
    for title, items in found:
        lines.extend([f"### {title}", "", *(f"- {i}" for i in items), ""])

    if not nlp.available:
        lines.extend([f"_{NO_SPACY}_", ""])
    if not found:
        lines.extend(["No issues found.", ""])
    return "\n".join(lines)


def build_docx(report: Report, nlp: NlpReport, source: str = "Pasted text") -> bytes:
    import docx

    document = docx.Document()
    document.add_heading(TITLE, level=0)
    document.add_paragraph(f"Source: {source}\nDate: {date.today().isoformat()}")

    document.add_heading("Stats", level=1)
    rows = stats_rows(report, nlp)
    table = document.add_table(rows=len(rows), cols=2)
    table.style = "Light List Accent 1"
    for row, (name, value) in zip(table.rows, rows):
        row.cells[0].text, row.cells[1].text = name, value

    document.add_heading("Suggestions", level=1)
    found = sections(report, nlp)
    for title, items in found:
        document.add_heading(title, level=2)
        for item in items:
            document.add_paragraph(item, style="List Bullet")
    if not nlp.available:
        document.add_paragraph(NO_SPACY).runs[0].italic = True
    if not found:
        document.add_paragraph("No issues found.")

    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


# DejaVu Sans (free licence, see fonts/LICENSE-DejaVu.txt) covers Latin, Greek, Cyrillic and the
# arrows and quotes the report uses. PDF's built-in fonts only cover Western European letters.
FONT_DIR = Path(__file__).parent / "fonts"
PDF_FONT = "DejaVu"
PDF_FONT_FILES = {"": "DejaVuSans.ttf", "B": "DejaVuSans-Bold.ttf", "I": "DejaVuSans-Oblique.ttf"}


def build_pdf(report: Report, nlp: NlpReport, source: str = "Pasted text") -> bytes:
    from fpdf import FPDF

    pdf = FPDF()
    pdf.set_title(TITLE)
    for style, file in PDF_FONT_FILES.items():
        pdf.add_font(PDF_FONT, style, str(FONT_DIR / file))
    pdf.add_page()
    covered = pdf.fonts[PDF_FONT.lower()].cmap  # fpdf2 keys fonts by lower-cased family + style

    def safe(text: str) -> str:
        # characters the font can't draw (e.g. Chinese or Hindi) would come out blank, so show "?"
        return "".join(c if ord(c) in covered or c in "\n\t" else "?" for c in text)

    def write(text: str, size: int = 11, style: str = "", gap: float = 1.5) -> None:
        pdf.set_font(PDF_FONT, style, size)
        pdf.multi_cell(0, size * 0.5, safe(text), new_x="LMARGIN", new_y="NEXT")
        pdf.ln(gap)

    write(TITLE, 20, "B", 2)
    write(f"Source: {source}\nDate: {date.today().isoformat()}", 10, gap=4)

    write("Stats", 15, "B")
    pdf.set_font(PDF_FONT, size=10)
    with pdf.table(col_widths=(60, 40), width=120, align="LEFT", first_row_as_headings=False) as table:
        for name, value in stats_rows(report, nlp):
            table.row([safe(name), safe(value)])
    pdf.ln(4)

    write("Suggestions", 15, "B")
    found = sections(report, nlp)
    for title, items in found:
        write(title, 12, "B", 1)
        for item in items:
            write(f"-  {item}", 10, gap=1)
        pdf.ln(2)
    if not nlp.available:
        write(NO_SPACY, 10, "I")
    if not found:
        write("No issues found.", 10)
    return bytes(pdf.output())
