import io

import docx
from pypdf import PdfReader

from analysis import analyze
from nlp_checks import NlpReport, check
from report import build_docx, build_markdown, build_pdf

MESSY = "We don't think outside the box. We analysed it and organized it. Due to the fact that it rained, we stayed."


def test_report_contains_stats_and_findings():
    text = "The report was written by the team. Due to the fact that it rained, we stayed."
    md = build_markdown(analyze(text), check(text), "essay.docx")
    assert "**Source:** essay.docx" in md
    assert "| Words |" in md
    assert "### Passive voice" in md
    assert "The report was written by the team." in md
    assert "due to the fact that" in md


def test_report_clean_text():
    text = "Cats chase mice."
    md = build_markdown(analyze(text), check(text))
    assert "No issues found." in md
    assert "###" not in md


def test_report_notes_missing_spacy():
    md = build_markdown(analyze("Cats chase mice."), NlpReport(available=False))
    assert "unavailable" in md


def test_report_has_new_checks_and_times():
    md = build_markdown(analyze(MESSY), NlpReport())
    assert "| Writing score |" in md and "| Reading time | under 1 min |" in md
    assert "### Clichés" in md and "think outside the box" in md
    assert "- don't ×1 → do not" in md
    assert "### Mixed UK/US spelling" in md and "UK: analysed ×1" in md


def test_docx_report():
    data = build_docx(analyze(MESSY), NlpReport(), "essay.pdf")
    document = docx.Document(io.BytesIO(data))
    text = "\n".join(p.text for p in document.paragraphs)
    assert "WriteLens report" in text and "Source: essay.pdf" in text
    assert "Clichés" in text and "don't ×1 → do not" in text
    assert document.tables[0].rows[0].cells[0].text == "Writing score"


def test_pdf_report_survives_unicode():
    data = build_pdf(analyze(MESSY + " Café “naïve” résumé — 漢字."), NlpReport(), "my essay 漢字.docx")
    text = "\n".join(page.extract_text() for page in PdfReader(io.BytesIO(data)).pages)
    assert "WriteLens report" in text and "my essay ??.docx" in text  # no Chinese in the font: shown as ?
    assert "don't ×1 → do not" in text  # the bundled font draws these
    assert "Writing score" in text
