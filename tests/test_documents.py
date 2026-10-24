import io

import docx
import pytest

from documents import DocumentError, extract_text


def make_docx(paragraphs: list[str], table: list[list[str]] | None = None) -> bytes:
    document = docx.Document()
    for p in paragraphs:
        document.add_paragraph(p)
    if table:
        t = document.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, value in enumerate(row):
                t.cell(r, c).text = value
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


def test_docx_paragraphs_and_tables():
    data = make_docx(["First paragraph.", "", "Second paragraph."], table=[["Cell text here.", "More."]])
    text = extract_text("essay.docx", data)
    assert "First paragraph." in text
    assert "Second paragraph." in text
    assert "Cell text here." in text


def make_pdf(lines: list[str], password: str = "") -> bytes:
    from fpdf import FPDF

    pdf = FPDF()
    if password:
        pdf.set_encryption(owner_password="owner", user_password=password)
    pdf.add_page()
    pdf.set_font("Helvetica", size=12)
    for line in lines:
        pdf.cell(0, 8, line, new_x="LMARGIN", new_y="NEXT")
    return bytes(pdf.output())


def test_pdf_text_and_rejoined_lines():
    data = make_pdf(["The first line of a para-", "graph keeps going here.", "Second sentence."])
    text = extract_text("essay.PDF", data)
    assert "The first line of a paragraph keeps going here." in text
    assert "Second sentence." in text


def test_pdf_with_empty_password_opens_and_real_password_is_refused():
    assert "Readable." in extract_text("a.pdf", make_pdf(["Readable."], password=""))
    with pytest.raises(DocumentError, match="password"):
        extract_text("locked.pdf", make_pdf(["Secret."], password="hunter2"))


def test_pdf_invalid_and_without_text():
    with pytest.raises(DocumentError, match="valid .pdf"):
        extract_text("broken.pdf", b"not a pdf at all")
    with pytest.raises(DocumentError, match="scan"):
        extract_text("scan.pdf", make_pdf([]))


def test_invalid_docx():
    with pytest.raises(DocumentError, match="valid .docx"):
        extract_text("broken.docx", b"not a zip file")


def test_tex_keeps_prose_drops_math_and_refs():
    src = r"""
\documentclass{article}
\usepackage{amsmath}
\begin{document}
\section{Introduction}
This is \textbf{important} work \cite{smith2020}. % a comment
The value $x^2$ grows, see Fig.~\ref{fig:one}.
\begin{equation}
E = mc^2
\end{equation}
We use 50\% of the data.
\end{document}
"""
    text = extract_text("paper.tex", src.encode())
    assert "Introduction" in text
    assert "This is important work" in text
    assert "50% of the data" in text
    for gone in ("smith2020", "a comment", "x^2", "mc^2", "fig:one", "\\", "amsmath"):
        assert gone not in text


def test_tex_heading_is_its_own_line():
    text = extract_text("p.tex", rb"\section{Methods}We measured it.")
    assert text.splitlines()[0] == "Methods"


def test_markdown_strips_markup():
    src = "# Title\n\nSome **bold** and [a link](http://x.com).\n\n```\ncode()\n```\n- item one\n"
    text = extract_text("notes.md", src.encode())
    assert "Title" in text
    assert "Some bold and a link." in text
    assert "item one" in text
    assert "code()" not in text and "http" not in text and "**" not in text


def test_txt_passthrough_and_bad_bytes():
    assert extract_text("a.txt", "Hello world.".encode()) == "Hello world."
    assert "Hello" in extract_text("b.txt", b"Hello \xff there.")


def test_unsupported_and_empty_and_too_large():
    with pytest.raises(DocumentError, match="Unsupported"):
        extract_text("image.png", b"...")
    with pytest.raises(DocumentError, match="No readable text"):
        extract_text("empty.txt", b"   \n  ")
    with pytest.raises(DocumentError, match="too large"):
        extract_text("big.txt", b"a" * (10 * 1024 * 1024 + 1))
